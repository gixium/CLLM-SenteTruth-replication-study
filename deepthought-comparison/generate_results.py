#!/usr/bin/env python3
"""
generate_results.py — DeepThought Results Aggregator & Report Generator

Aggregates raw simulation CSV files into summary statistics, Excel reports,
comparison charts (PNG), and LaTeX-ready tables.

Usage:
    python3 generate_results.py [--results-dir results]

Author: Research Project — CLLM-SenteTruth Replication Study
"""

import argparse
import csv
import glob
import os
import re
import sys
import subprocess
from collections import defaultdict

# ---------------------------------------------------------------------------
# Check dependencies
# ---------------------------------------------------------------------------
try:
    import numpy as np
except ImportError:
    print("ERROR: numpy not installed. Run: pip install numpy")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    print("WARNING: openpyxl not installed. Excel reports will be skipped.")
    print("         Install with: pip install openpyxl")

try:
    import matplotlib
    matplotlib.use("Agg")  # non-interactive backend
    import matplotlib.pyplot as plt
    HAS_MPL = True
except ImportError:
    HAS_MPL = False
    print("WARNING: matplotlib not installed. Charts will be skipped.")
    print("         Install with: pip install matplotlib")

# CLLM Accuracy Data (for q* analysis)

Q_MIN_RATIONAL = 0.50
Q_MIN_PAPER = 0.80
ACC_ROBUST_THRESHOLD = 95.0

CLLM_TABLE = [
    # 60/40 split (k=4 adversarial out of 10)
    ("GPT-4o-mini", "MIX", "60-40", 4, 100.0, 67.0,  1.0, 99.0),
    ("GPT-4o-mini", "PRO", "60-40", 4, 100.0,  3.3,  1.7, 100.0),
    ("Gemini",      "MIX", "60-40", 4, 100.0, 58.0,  8.0, 100.0),
    ("Gemini",      "PRO", "60-40", 4, 100.0, 28.3, 11.7, 100.0),
    ("DS-Chat",     "MIX", "60-40", 4,  50.0,  1.0,  1.0,  3.0),
    ("DS-Chat",     "PRO", "60-40", 4,  40.0,  0.0,  0.0,  0.0),
    ("DS-Reasoner", "MIX", "60-40", 4,   0.0,  2.0,  0.0,  1.0),
    ("DS-Reasoner", "PRO", "60-40", 4,  15.0,  8.3,  5.0,  6.7),
    # 70/30 split (k=3 adversarial out of 10)
    ("GPT-4o-mini", "MIX", "70-30", 3, 100.0, 100.0, 100.0, 100.0),
    ("GPT-4o-mini", "PRO", "70-30", 3, 100.0, 100.0, 56.7, 100.0),
    ("Gemini",      "MIX", "70-30", 3, 100.0, 100.0, 100.0, 100.0),
    ("Gemini",      "PRO", "70-30", 3, 100.0, 100.0, 100.0, 100.0),
    ("DS-Chat",     "MIX", "70-30", 3, 100.0, 67.0, 43.0, 53.0),
    ("DS-Chat",     "PRO", "70-30", 3, 100.0, 13.3,  0.0, 61.7),
    ("DS-Reasoner", "MIX", "70-30", 3,  17.0,  8.0, 59.0, 77.0),
    ("DS-Reasoner", "PRO", "70-30", 3, 100.0, 100.0, 53.3, 60.0),
]

CLLM_CONFIGS = ["C1", "C2", "C3", "C4"]



# Data Loading

def load_raw_csvs(raw_dir: str) -> list:
    """Load all raw DT CSV files from the raw directory."""
    pattern = os.path.join(raw_dir, "dt_*.csv")
    files = sorted(glob.glob(pattern))
    if not files:
        print(f"  No CSV files found matching: {pattern}")
        return []

    all_rows = []
    for fpath in files:
        basename = os.path.basename(fpath)
        with open(fpath, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                row["_source_file"] = basename
                all_rows.append(row)

    print(f"  Loaded {len(all_rows)} rows from {len(files)} CSV files.")
    return all_rows


def load_reputation_logs(raw_dir: str, dataset: str, split: str,
                          accuracy: float, n_runs: int) -> list:
    """Load reputation log files for a given config."""
    logs = []
    for run_idx in range(1, n_runs + 1):
        fname = f"reputation_log_{dataset}_{split}_acc{accuracy}_run{run_idx}.txt"
        fpath = os.path.join(raw_dir, fname)
        if os.path.exists(fpath):
            run_log = []
            with open(fpath, "r") as f:
                for line in f:
                    vals = list(map(int, line.strip().split()))
                    run_log.append(vals)
            logs.append(run_log)
    return logs


# Aggregation

def aggregate_results(rows: list) -> list:
    """
    Group by (dataset, split, accuracy) and compute summary statistics.
    Returns a list of summary dicts.
    """
    config_groups = defaultdict(list)
    for row in rows:
        fname = row.get("_source_file", "")
        m = re.match(r"dt_(\w+)_(\d+-\d+)_acc([\d.]+)\.csv", fname)
        if m:
            key = (m.group(1), m.group(2), m.group(3))
        else:
            key = (row.get("accuracy", "?"), row.get("adv_control", "?"), fname)
        config_groups[key].append(row)

    summaries = []
    for (dataset, split, acc_str), group_rows in sorted(config_groups.items()):
        corrupted_vals = [int(r["prop_corrupted"]) for r in group_rows]
        n_props = int(group_rows[0]["propositions"])
        n_voters = int(group_rows[0]["voters"])
        adv_control = float(group_rows[0]["adv_control"])
        accuracy = float(acc_str)

        target_vals = [int(r["target_corrupted"]) for r in group_rows]

        avg_c = np.mean(corrupted_vals)
        std_c = np.std(corrupted_vals, ddof=1) if len(corrupted_vals) > 1 else 0.0

        summaries.append({
            "dataset": dataset,
            "split": split,
            "accuracy": accuracy,
            "adv_control": adv_control,
            "n_voters": n_voters,
            "n_propositions": n_props,
            "n_runs": len(group_rows),
            "corruption_avg": round(float(avg_c), 2),
            "corruption_std": round(float(std_c), 2),
            "corruption_min": min(corrupted_vals),
            "corruption_max": max(corrupted_vals),
            "system_accuracy": round((1.0 - float(avg_c) / n_props) * 100, 2),
            "target_corrupted_pct": round(
                sum(target_vals) / len(target_vals) * 100, 2
            ) if target_vals else 0.0,
        })

    return summaries


# Summary CSVs

def write_summary_csvs(summaries: list, results_dir: str):
    """Write summary CSV files grouped by split."""
    splits = set(s["split"] for s in summaries)

    header = [
        "dataset", "split", "accuracy", "adv_control", "n_voters",
        "n_propositions", "n_runs", "corruption_avg", "corruption_std",
        "corruption_min", "corruption_max", "system_accuracy",
        "target_corrupted_pct",
    ]

    for split in sorted(splits):
        split_rows = [s for s in summaries if s["split"] == split]
        fpath = os.path.join(results_dir, f"dt_summary_{split}.csv")
        with open(fpath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=header)
            writer.writeheader()
            for row in sorted(split_rows,
                               key=lambda x: (x["dataset"], x["accuracy"])):
                writer.writerow({k: row[k] for k in header})
        print(f"  Summary CSV -> {fpath}")


# Excel Report

def write_excel_report(summaries: list, all_rows: list, results_dir: str):
    """Generate a formatted Excel report."""
    if openpyxl is None:
        print("  Skipping Excel report (openpyxl not installed).")
        return

    fpath = os.path.join(results_dir, "dt_comparison_report.xlsx")
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496",
                              fill_type="solid")
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE",
                             fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE",
                           fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C",
                              fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    def style_header(ws, n_cols):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

    def auto_width(ws):
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = (
                min(max_len + 3, 25)
            )

    for split in sorted(set(s["split"] for s in summaries)):
        ws = wb.create_sheet(title=f"Summary_{split}")
        headers = [
            "Dataset", "Accuracy", "Adv%", "Voters", "Props", "Runs",
            "C-AVG", "STD", "MIN", "MAX", "Sys.Accuracy%", "Target%",
        ]
        ws.append(headers)
        style_header(ws, len(headers))

        for s in sorted(
            [x for x in summaries if x["split"] == split],
            key=lambda x: (x["dataset"], x["accuracy"]),
        ):
            row = [
                s["dataset"], s["accuracy"], s["adv_control"], s["n_voters"],
                s["n_propositions"], s["n_runs"], s["corruption_avg"],
                s["corruption_std"], s["corruption_min"], s["corruption_max"],
                s["system_accuracy"], s["target_corrupted_pct"],
            ]
            ws.append(row)

            r = ws.max_row
            acc_cell = ws.cell(row=r, column=11)
            if s["system_accuracy"] >= 95:
                acc_cell.fill = green_fill
            elif s["system_accuracy"] >= 80:
                acc_cell.fill = yellow_fill
            else:
                acc_cell.fill = red_fill

            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).alignment = center_align
                ws.cell(row=r, column=col).border = thin_border

        auto_width(ws)

    ws = wb.create_sheet(title="All_Runs")
    if all_rows:
        run_headers = [
            "run", "voters", "propositions", "accuracy", "adv_control",
            "prop_corrupted", "prop_corrupted_pct", "target_corrupted",
            "alpha", "stake", "elapsed_time", "source_file",
        ]
        ws.append(run_headers)
        style_header(ws, len(run_headers))

        for row in all_rows:
            ws.append([
                row.get("run", ""), row.get("voters", ""),
                row.get("propositions", ""), row.get("accuracy", ""),
                row.get("adv_control", ""), row.get("prop_corrupted", ""),
                row.get("prop_corrupted_pct", ""), row.get("target_corrupted", ""),
                row.get("alpha", ""), row.get("stake", ""),
                row.get("elapsed_time", ""), row.get("_source_file", ""),
            ])
        auto_width(ws)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(fpath)
    print(f"  Excel -> {fpath}")


# Charts

def generate_charts(summaries: list, raw_dir: str, results_dir: str):
    """Generate comparison charts as PNG files."""
    if not HAS_MPL:
        print("  Skipping charts (matplotlib not installed).")
        return

    charts_dir = os.path.join(results_dir, "charts")
    os.makedirs(charts_dir, exist_ok=True)

    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.size": 10,
        "axes.grid": True,
        "grid.alpha": 0.3,
        "figure.dpi": 300,
    })

    colors = {"MIX": "#2F5496", "PRO": "#C00000"}

    # Chart 1: System Accuracy by voter accuracy (per split)
    for split in sorted(set(s["split"] for s in summaries)):
        fig, ax = plt.subplots(figsize=(7, 4.5))

        for dataset in ["MIX", "PRO"]:
            subset = sorted(
                [s for s in summaries
                 if s["split"] == split and s["dataset"] == dataset],
                key=lambda x: x["accuracy"],
            )
            if not subset:
                continue

            accs = [s["accuracy"] for s in subset]
            sys_accs = [s["system_accuracy"] for s in subset]
            stds = [s["corruption_std"] / s["n_propositions"] * 100
                    for s in subset]

            ax.errorbar(accs, sys_accs, yerr=stds, marker="o", markersize=7,
                        linewidth=2, label=dataset,
                        color=colors.get(dataset, "#333"), capsize=4,
                        capthick=1.5)

        ax.set_xlabel("Honest Voter Accuracy", fontsize=11)
        ax.set_ylabel("System Accuracy (%)", fontsize=11)
        ax.set_title(f"DeepThought - System Accuracy vs Voter Accuracy "
                     f"(Split {split})", fontsize=12, fontweight="bold")
        ax.legend(fontsize=10)
        ax.set_ylim(0, 105)
        ax.set_xlim(0.45, 1.0)

        fpath = os.path.join(charts_dir, f"accuracy_comparison_{split}.png")
        fig.tight_layout()
        fig.savefig(fpath, dpi=300, bbox_inches="tight")
        plt.close(fig)
        print(f"  Chart -> {fpath}")

    # Chart 2: Corruption rate by accuracy (per split)
    for split in sorted(set(s["split"] for s in summaries)):
        fig, ax = plt.subplots(figsize=(7, 4.5))

        for dataset in ["MIX", "PRO"]:
            subset = sorted(
                [s for s in summaries
                 if s["split"] == split and s["dataset"] == dataset],
                key=lambda x: x["accuracy"],
            )
            if not subset:
                continue

            accs = [s["accuracy"] for s in subset]
            corr_avgs = [s["corruption_avg"] for s in subset]
            corr_stds = [s["corruption_std"] for s in subset]

            ax.errorbar(accs, corr_avgs, yerr=corr_stds, marker="s",
                        markersize=7, linewidth=2, label=dataset,
                        color=colors.get(dataset, "#333"), capsize=4,
                        capthick=1.5)

        ax.set_xlabel("Honest Voter Accuracy", fontsize=11)
        ax.set_ylabel("Average Corrupted Propositions", fontsize=11)
        ax.set_title(f"DeepThought - Corruption Rate vs Voter Accuracy "
                     f"(Split {split})", fontsize=12, fontweight="bold")
        ax.legend(fontsize=10)
        ax.set_xlim(0.45, 1.0)

        fpath = os.path.join(charts_dir, f"corruption_by_accuracy_{split}.png")
        fig.tight_layout()
        fig.savefig(fpath, dpi=300, bbox_inches="tight")
        plt.close(fig)
        print(f"  Chart -> {fpath}")

    # Chart 3: Reputation evolution
    for s in summaries:
        dataset, split, accuracy = s["dataset"], s["split"], s["accuracy"]
        n_runs = s["n_runs"]
        n_honest = 6 if split == "60-40" else 7
        n_adversarial = 10 - n_honest

        logs = load_reputation_logs(raw_dir, dataset, split, accuracy, n_runs)
        if not logs:
            continue

        n_steps = min(len(log) for log in logs)
        avg_honest_rep = np.zeros(n_steps)
        avg_adv_rep = np.zeros(n_steps)

        for log in logs:
            for step in range(n_steps):
                avg_honest_rep[step] += np.mean(log[step][:n_honest])
                avg_adv_rep[step] += np.mean(log[step][n_honest:])

        avg_honest_rep /= len(logs)
        avg_adv_rep /= len(logs)

        fig, ax = plt.subplots(figsize=(7, 4))
        steps = range(1, n_steps + 1)
        ax.plot(steps, avg_honest_rep, linewidth=2, color="#2F5496",
                label=f"Honest voters (n={n_honest})")
        ax.plot(steps, avg_adv_rep, linewidth=2, color="#C00000",
                label=f"Adversarial voters (n={n_adversarial})")

        ax.set_xlabel("Proposition Step", fontsize=11)
        ax.set_ylabel("Average Reputation", fontsize=11)
        ax.set_title(f"Reputation Evolution - {dataset} {split} "
                     f"(acc={accuracy})", fontsize=12, fontweight="bold")
        ax.legend(fontsize=10)

        fpath = os.path.join(
            charts_dir,
            f"reputation_evolution_{dataset}_{split}_acc{accuracy}.png",
        )
        fig.tight_layout()
        fig.savefig(fpath, dpi=300, bbox_inches="tight")
        plt.close(fig)
        print(f"  Chart -> {fpath}")


# LaTeX Tables

def write_latex_tables(summaries: list, results_dir: str):
    """Generate LaTeX-ready tables for the paper."""
    fpath = os.path.join(results_dir, "dt_latex_tables.tex")

    with open(fpath, "w", encoding="utf-8") as f:
        f.write("% DeepThought Simulation Results - Auto-generated\n")
        f.write("% Include this file in your LaTeX document.\n\n")

        for split in sorted(set(s["split"] for s in summaries)):
            subset = sorted(
                [s for s in summaries if s["split"] == split],
                key=lambda x: (x["dataset"], x["accuracy"]),
            )
            if not subset:
                continue

            adv_pct = int(float(subset[0]["adv_control"]) * 100)
            n_voters = subset[0]["n_voters"]
            n_runs = subset[0]["n_runs"]

            f.write(f"% Split: {split} ({adv_pct}% adversarial)\n")
            f.write("\\begin{table}[h]\n")
            f.write("\\centering\n")
            f.write(f"\\caption{{DeepThought simulation results "
                    f"(Split {split}, {n_voters} voters, "
                    f"{n_runs} repetitions)}}\n")
            f.write("\\label{tab:dt-" + split + "}\n")
            f.write("\\begin{tabular}{llrrrrr}\n")
            f.write("\\toprule\n")
            f.write("Dataset & Voter Acc. & C-AVG & STD & MIN & MAX "
                    "& Sys. Acc. \\\\\n")
            f.write("\\midrule\n")

            for s in subset:
                f.write(
                    f"{s['dataset']} & {s['accuracy']:.2f} & "
                    f"{s['corruption_avg']:.2f} & {s['corruption_std']:.2f} & "
                    f"{s['corruption_min']} & {s['corruption_max']} & "
                    f"{s['system_accuracy']:.2f}\\% \\\\\n"
                )

            f.write("\\bottomrule\n")
            f.write("\\end{tabular}\n")
            f.write("\\end{table}\n\n")

    print(f"  LaTeX -> {fpath}")


# Empirical q* Analysis

def build_accuracy_curves(summaries: list) -> dict:
    """Build empirical curves from the summary data."""
    curves = {}
    for s in summaries:
        key = (s["dataset"], s["split"])
        if key not in curves:
            curves[key] = []
        curves[key].append((s["accuracy"], s["system_accuracy"]))
    
    # Sort points for each curve by q
    for key in curves:
        curves[key].sort(key=lambda x: x[0])
    
    return curves

def find_q_star_from_curve(target_acc: float, curve: list) -> float:
    """Find q* by linear interpolation on the empirical curve."""
    if not curve:
        return None
        
    if target_acc <= 0.0:
        return 0.0
        
    # Check bounds
    if target_acc <= curve[0][1]:
        return curve[0][0]
    if target_acc >= curve[-1][1]:
        return None # Unreachable
        
    for i in range(len(curve) - 1):
        q1, acc1 = curve[i]
        q2, acc2 = curve[i+1]
        
        if acc1 <= target_acc <= acc2:
            if acc1 == acc2:
                return q1
            # Linear interpolation
            slope = (q2 - q1) / (acc2 - acc1)
            return q1 + slope * (target_acc - acc1)
            
    return None

def run_q_star_analysis(summaries: list, results_dir: str):
    """Run full empirical q* analysis."""
    print("\n" + "=" * 60)
    print("  DeepThought Empirical q* Analysis")
    print("=" * 60)
    
    curves = build_accuracy_curves(summaries)
    all_results = []
    
    for model, ds, split, k, *accs in CLLM_TABLE:
        for cfg, acc in zip(CLLM_CONFIGS, accs):
            vulnerable = acc < ACC_ROBUST_THRESHOLD
            
            if not vulnerable:
                all_results.append({
                    "split": split, "model": model, "ds": ds, "cfg": cfg,
                    "k": k, "acc": acc, "q_star": None, "vulnerable": False
                })
                continue
                
            curve = curves.get((ds, split), [])
            q_star = find_q_star_from_curve(acc, curve)
            
            all_results.append({
                "split": split, "model": model, "ds": ds, "cfg": cfg,
                "k": k, "acc": acc, "q_star": q_star, "vulnerable": True
            })

    # Generate Report
    fpath = os.path.join(results_dir, "q_star_deepthought_output.txt")
    with open(fpath, "w", encoding="utf-8") as f:
        f.write(f"DeepThought q* (empirical interpolation) | N=10\n")
        f.write("-" * 85 + "\n")
        f.write(f"{'Split':7} {'Model':13} {'Ds':4} {'Cfg':4} {'Acc%':>6} {'q*':>7}  Note\n")
        f.write("-" * 85 + "\n")
        
        for r in all_results:
            marker = "[V]" if r["vulnerable"] else "   "
            if not r["vulnerable"]:
                note = "C-LLM unreachable"
                q_str = "N/A"
            elif r["q_star"] is None:
                note = "DT unreachable"
                q_str = "N/A"
            elif r["q_star"] <= 0.001:
                note = "DT wins (q* ~ 0 = trivial)"
                q_str = "0.000"
            elif r["q_star"] <= Q_MIN_RATIONAL:
                note = f"DT wins (q* <= {Q_MIN_RATIONAL} = rational minimum)"
                q_str = f"{r['q_star']:.3f}"
            elif r["q_star"] <= Q_MIN_PAPER:
                note = f"DT wins (q* <= {Q_MIN_PAPER} = paper standard)"
                q_str = f"{r['q_star']:.3f}"
            else:
                note = f"DT wins if q* >= {r['q_star']:.3f}"
                q_str = f"{r['q_star']:.3f}"
                
            f.write(f"{r['split']:7} {r['model']:13} {r['ds']:4} {r['cfg']:4} {r['acc']:>5.1f}% {marker} {q_str:>7}  {note}\n")
    
    print(f"  q* Output -> {fpath}")
    
    # Save unique q* values for verification
    q_stars = set()
    for r in all_results:
        if r["q_star"] is not None and r["q_star"] > 0.001:
            q_stars.add((r["ds"], r["split"], round(r["q_star"], 3)))
            
    q_star_file = os.path.join(results_dir, "q_star_values_to_verify.csv")
    with open(q_star_file, "w") as f:
        f.write("dataset,split,accuracy\n")
        for ds, split, q in sorted(list(q_stars)):
            f.write(f"{ds},{split},{q}\n")
            
    print(f"  Unique q* to verify -> {q_star_file}")


def run_q_star_verification(results_dir: str):
    """Run verification experiments for determined q* values."""
    print("\n" + "=" * 60)
    print("  DeepThought q* Verification")
    print("=" * 60)
    
    q_star_file = os.path.join(results_dir, "q_star_values_to_verify.csv")
    if not os.path.exists(q_star_file):
        print(f"  Error: {q_star_file} not found. Run analysis first.")
        return
        
    sim_script = os.path.join(os.path.dirname(__file__), "deepthought_sim.py")
    out_dir = os.path.join(results_dir, "raw_verification")
    os.makedirs(out_dir, exist_ok=True)
    
    with open(q_star_file, "r") as f:
        lines = f.readlines()[1:] # skip header
        
    for idx, line in enumerate(lines):
        ds, split, q_str = line.strip().split(",")
        q = float(q_str)
        print(f"  [{idx+1}/{len(lines)}] Verifying: {ds} {split} q*={q}")
        cmd = [
            "python3", sim_script,
            "--dataset", ds, "--split", split, "--accuracy", str(q),
            "--repetitions", "30", "--output-dir", out_dir
        ]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL)
        
    print(f"  Verification complete. Raw data saved to {out_dir}")


# Main

def main():
    parser = argparse.ArgumentParser(
        description="DeepThought Results Aggregator & Report Generator",
    )
    parser.add_argument(
        "--results-dir", type=str, default="results",
        help="Results directory containing raw/ subdirectory. Default: results",
    )
    parser.add_argument(
        "--q-star-analysis", action="store_true",
        help="Run empirical q* interpolation from curve data",
    )
    parser.add_argument(
        "--verify-q-star", action="store_true",
        help="Run verification experiments for the found q* values",
    )
    args = parser.parse_args()

    results_dir = args.results_dir
    raw_dir = os.path.join(results_dir, "raw")

    if args.verify_q_star:
        run_q_star_verification(results_dir)
        return

    print(f"\n{'=' * 60}")
    print(f"  DeepThought Results Generator")
    print(f"{'=' * 60}")

    print(f"\n  Loading raw CSV files from: {raw_dir}")
    all_rows = load_raw_csvs(raw_dir)
    if not all_rows:
        print("  No data found. Run deepthought_sim.py first.")
        sys.exit(1)

    print(f"\n  Aggregating results...")
    summaries = aggregate_results(all_rows)
    print(f"  {len(summaries)} configurations found.")

    print(f"\n  Writing summary CSVs...")
    write_summary_csvs(summaries, results_dir)

    print(f"\n  Generating Excel report...")
    write_excel_report(summaries, all_rows, results_dir)

    print(f"\n  Generating charts...")
    generate_charts(summaries, raw_dir, results_dir)

    print(f"\n  Writing LaTeX tables...")
    write_latex_tables(summaries, results_dir)

    if args.q_star_analysis:
        run_q_star_analysis(summaries, results_dir)

    print(f"\n{'=' * 60}")
    print(f"  Done! All outputs in: {results_dir}/")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    main()
